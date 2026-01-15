import os
import PyPDF2
from pdf2docx import Converter
import pytesseract
from PIL import Image
from pdf2image import convert_from_path
from pathlib import Path
from typing import Tuple
from app.core.exceptions import (
    InvalidFileError, 
    PasswordProtectedError, 
    OCRDetectionError, 
    ConversionError
)

class PDFService:
    """Service for PDF analysis and conversion operations"""
    
    @staticmethod
    def detect_scanned_pdf(pdf_path: str) -> bool:
        """
        Detect if PDF is scanned (image-based) or text-based.
        Raises specific exceptions for invalid files.
        """
        try:
            with open(pdf_path, 'rb') as file:
                try:
                    reader = PyPDF2.PdfReader(file)
                    
                    if reader.is_encrypted:
                        raise PasswordProtectedError("Cannot process encrypted PDF. Please remove password.")
                    
                    # Check first 3 pages
                    pages_to_check = min(3, len(reader.pages))
                    
                    # Check for empty PDF
                    if len(reader.pages) == 0:
                        raise InvalidFileError("PDF file is empty")

                    for i in range(pages_to_check):
                        page = reader.pages[i]
                        text = page.extract_text()
                        if text and len(text.strip()) > 50:
                            return False # Found meaningful text
                            
                    return True # No text found, likely scanned
                    
                except PyPDF2.errors.PdfReadError:
                    raise InvalidFileError("File is corrupted or not a valid PDF")
                    
        except (PasswordProtectedError, InvalidFileError):
            raise
        except Exception as e:
            print(f"Error detecting PDF type: {e}")
            # If we can't read it properly, we shouldn't guess
            raise InvalidFileError(f"Could not analyze PDF: {str(e)}")
    
    @staticmethod
    def convert_text_based_pdf(pdf_path: str, output_path: str) -> None:
        """Convert text-based PDF using pdf2docx"""
        try:
            cv = Converter(pdf_path)
            cv.convert(output_path)
            cv.close()
        except Exception as e:
            msg = str(e).lower()
            if "encrypted" in msg:
                raise PasswordProtectedError()
            raise ConversionError(f"Text conversion failed: {str(e)}")
    
    @staticmethod
    def convert_scanned_pdf_with_ocr(pdf_path: str, output_path: str) -> None:
        """
        Convert scanned PDF using OCR using a 'Dual Output' strategy.
        To ensure no data is lost (e.g. passport photos, complex tables) and text is editable:
        1. Insert the Page Image (Visual fidelity).
        2. Insert the Extracted Text below it (Editability).
        """
        import tempfile
        try:
            from docx import Document
            from docx.shared import Inches, Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            
            # Convert PDF pages to images
            try:
                images = convert_from_path(pdf_path)
            except Exception as e:
                raise OCRDetectionError(f"Failed to convert PDF to images: {str(e)}")

            if not images:
                raise InvalidFileError("No pages found in PDF")
            
            doc = Document()
            
            # Create a temporary directory for processing images
            with tempfile.TemporaryDirectory() as temp_dir:
                for i, image in enumerate(images):
                    print(f"Processing page {i+1}/{len(images)} with OCR...")
                    
                    # 1. OCR Extraction
                    try:
                        text = pytesseract.image_to_string(image)
                    except Exception as e:
                        print(f"OCR failed for page {i+1}: {e}")
                        text = ""
                    
                    # 2. Insert Page Image (Preserves photos, layout, tables)
                    img_path = os.path.join(temp_dir, f"page_{i}.png")
                    image.save(img_path, "PNG")
                    
                    try:
                        # Add image centered
                        doc.add_picture(img_path, width=Inches(6.0))
                        last_paragraph = doc.paragraphs[-1] 
                        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    except Exception as e:
                        print(f"Failed to insert image for page {i+1}: {e}")

                    # 3. Insert Extracted Text (Provides editability)
                    if text.strip():
                        # Add a separator
                        p = doc.add_paragraph()
                        run = p.add_run("--- Extracted Text Below ---")
                        run.font.size = Pt(8)
                        run.font.italic = True
                        run.font.color.rgb = None # Default color
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        
                        doc.add_paragraph(text)
                    else:
                        # Inform user if no text found
                        p = doc.add_paragraph()
                        run = p.add_run("[No text text detected on this page - Image only]")
                        run.font.size = Pt(8)
                        run.font.italic = True
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                    # Add page break except for last page
                    if i < len(images) - 1:
                        doc.add_page_break()
            
            doc.save(output_path)
            
        except (OCRDetectionError, InvalidFileError):
            raise
        except Exception as e:
            raise OCRDetectionError(f"OCR process failed: {str(e)}")
    
    @staticmethod
    def convert_pdf_to_docx(pdf_path: str, output_path: str) -> Tuple[bool, bool]:
        """
        Main conversion method with automatic fallback.
        Returns: (success, used_ocr)
        Now handles exceptions gracefully or re-raises them.
        """
        try:
            # First, validation & detection
            is_scanned = PDFService.detect_scanned_pdf(pdf_path)
            
            if is_scanned:
                print("Detected scanned PDF - using OCR conversion")
                PDFService.convert_scanned_pdf_with_ocr(pdf_path, output_path)
                return True, True
            else:
                print("Detected text-based PDF - using direct conversion")
                try:
                    PDFService.convert_text_based_pdf(pdf_path, output_path)
                    return True, False
                except Exception as e:
                    print(f"Text-based conversion failed ({e}) - attempting OCR fallback")
                    # Fallback to OCR
                    PDFService.convert_scanned_pdf_with_ocr(pdf_path, output_path)
                    return True, True
                    
        except PDFConverterException as e:
            # Re-raise known exceptions for the caller to handle
            raise e
        except Exception as e:
            # Wrap unknown errors
            raise ConversionError(f"Unexpected error: {str(e)}")

    @staticmethod
    def validate_pdf_file(file_path: str) -> bool:
        """Strict validation of PDF file"""
        if not os.path.exists(file_path):
            return False
        try:
            # Check if file is empty
            if os.path.getsize(file_path) == 0:
                raise InvalidFileError("File is 0 bytes")

            # Check PDF header
            with open(file_path, 'rb') as f:
                header = f.read(5)
                if header != b'%PDF-':
                    raise InvalidFileError("Invalid PDF header")

            # Try parsing with PyPDF2
            with open(file_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                if reader.is_encrypted:
                    raise PasswordProtectedError()
                # Try accessing pages
                _ = len(reader.pages)
                
            return True
            
        except PDFConverterException:
            raise
        except Exception as e:
            print(f"PDF validation failed: {e}")
            raise InvalidFileError(f"File validation failed: {str(e)}")

